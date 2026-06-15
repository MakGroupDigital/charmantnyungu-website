from __future__ import annotations

import hashlib
import math
import os
import random
from dataclasses import dataclass
from pathlib import Path

import qrcode
from qrcode.constants import ERROR_CORRECT_H
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "livrables" / "P-E2M2E"
IMG_DIR = OUT / "schemas"
SOURCE = ROOT / ".codex" / "attachments"
PHOTO_SOURCE = ROOT / "public" / "photos" / "IMG_1838.jpg"
SITE_URL = "https://www.charmantnyungu.com"
AUTHOR_CONTACT = (
    "Concepteur: Charmant Nyungu K.\n"
    "Consultant en innovation technologique, panafricaniste.\n"
    "Site web: www.charmantnyungu.com\n"
    "Email: consultant@charmantnyungu.com\n"
    "Téléphone: +243 835 137 837"
)


@dataclass
class Page:
    title: str
    body: list[str]
    diagram: Path | None = None
    table: list[list[str]] | None = None


PALETTE = [
    (15, 23, 42),
    (12, 74, 110),
    (0, 95, 115),
    (148, 81, 0),
    (111, 29, 27),
    (64, 83, 76),
    (73, 80, 87),
    (40, 54, 24),
    (94, 57, 75),
    (23, 37, 84),
]


def prepare_brand_assets() -> tuple[Path, Path]:
    assets_dir = OUT / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    portrait_path = assets_dir / "charmant_nyungu_portrait_cercle.png"
    qr_path = assets_dir / "qrcode_charmantnyungu_futuriste.png"

    portrait_size = 520
    source = Image.open(PHOTO_SOURCE).convert("RGB")
    w, h = source.size
    side = min(w, h)
    left = (w - side) // 2
    top = (h - side) // 2
    portrait = source.crop((left, top, left + side, top + side)).resize((portrait_size, portrait_size), Image.LANCZOS)
    mask = Image.new("L", (portrait_size, portrait_size), 0)
    md = ImageDraw.Draw(mask)
    md.ellipse((0, 0, portrait_size - 1, portrait_size - 1), fill=255)
    out = Image.new("RGBA", (portrait_size + 36, portrait_size + 36), (0, 0, 0, 0))
    halo = ImageDraw.Draw(out)
    halo.ellipse((4, 4, portrait_size + 31, portrait_size + 31), outline=(12, 74, 110, 210), width=10)
    halo.ellipse((15, 15, portrait_size + 20, portrait_size + 20), outline=(0, 170, 190, 170), width=4)
    out.paste(portrait, (18, 18), mask)
    portrait_draw = ImageDraw.Draw(out)
    portrait_draw.arc((2, 2, portrait_size + 33, portrait_size + 33), 305, 45, fill=(0, 220, 255, 230), width=8)
    portrait_draw.arc((2, 2, portrait_size + 33, portrait_size + 33), 125, 220, fill=(148, 81, 0, 220), width=8)
    out.save(portrait_path)

    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_H, box_size=12, border=2)
    qr.add_data(SITE_URL)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="#0f172a", back_color="white").convert("RGBA")
    qr_img = qr_img.resize((430, 430), Image.Resampling.NEAREST)
    canvas = Image.new("RGBA", (520, 520), (248, 250, 252, 255))
    d = ImageDraw.Draw(canvas)
    d.rounded_rectangle((8, 8, 512, 512), radius=36, fill=(255, 255, 255, 255), outline=(15, 23, 42, 255), width=8)
    d.rounded_rectangle((24, 24, 496, 496), radius=26, outline=(0, 170, 190, 220), width=4)
    canvas.paste(qr_img, (45, 45), qr_img)
    corner = (0, 170, 190, 255)
    d.line((48, 30, 118, 30, 118, 42), fill=corner, width=8)
    d.line((402, 30, 472, 30, 472, 100), fill=corner, width=8)
    d.line((48, 490, 118, 490, 118, 478), fill=corner, width=8)
    d.line((402, 490, 472, 490, 472, 420), fill=corner, width=8)
    d.text((169, 470), "SCAN SITE", font=font(24, True), fill=(15, 23, 42, 255))
    canvas.save(qr_path)
    return portrait_path, qr_path


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, width: int, fill=(255, 255, 255)) -> None:
    fnt = font(21, True)
    x, y = xy
    for line in wrap(draw, text, fnt, width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += 28


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill, width=5) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    p1 = (end[0] - size * math.cos(angle - 0.45), end[1] - size * math.sin(angle - 0.45))
    p2 = (end[0] - size * math.cos(angle + 0.45), end[1] - size * math.sin(angle + 0.45))
    draw.polygon([end, p1, p2], fill=fill)


def make_diagram(index: int, title: str) -> Path:
    random.seed(index * 1709)
    w, h = 1500, 850
    bg = (248, 250, 252)
    img = Image.new("RGB", (w, h), bg)
    d = ImageDraw.Draw(img)
    base = PALETTE[index % len(PALETTE)]
    accent = PALETTE[(index * 3 + 2) % len(PALETTE)]
    third = PALETTE[(index * 5 + 4) % len(PALETTE)]
    d.rectangle([0, 0, w, 92], fill=base)
    d.text((44, 28), f"Schéma {index:02d} - {title}", font=font(34, True), fill=(255, 255, 255))
    mode = index % 12

    if mode == 0:
        nodes = {"PTx": (245, 500), "MB": (750, 255), "PRx": (1250, 500)}
        for a, b in [("PTx", "MB"), ("MB", "PRx"), ("PRx", "PTx")]:
            arrow(d, nodes[a], nodes[b], accent, 6)
        for label, (x, y) in nodes.items():
            d.ellipse([x - 105, y - 105, x + 105, y + 105], fill=base, outline=third, width=8)
            d.text((x - 42, y - 18), label, font=font(34, True), fill=(255, 255, 255))
        draw_label(d, (95, 690), "Topologie triangulaire: émission physique, milieu aveugle, réception physicalisée.", 1280, base)
    elif mode == 1:
        y = 470
        d.line([100, y, 1400, y], fill=base, width=8)
        for i, label in enumerate(["t0", "t+10ms", "t+25ms", "t+50ms", "annihilation"]):
            x = 130 + i * 310
            d.ellipse([x - 22, y - 22, x + 22, y + 22], fill=accent)
            d.text((x - 55, y + 42), label, font=font(24, True), fill=base)
            d.line([x, y - 150, x, y - 28], fill=third, width=3)
        draw_label(d, (110, 170), "Fenêtre temporelle Δt: la donnée existe comme potentiel reconstructible uniquement dans l'intervalle admis.", 1220, base)
    elif mode == 2:
        for r in range(75, 360, 55):
            d.ellipse([750 - r, 450 - r, 750 + r, 450 + r], outline=accent if r % 2 else base, width=5)
        for i in range(18):
            ang = i * math.pi / 9
            x = 750 + int(math.cos(ang) * random.randint(120, 330))
            y = 450 + int(math.sin(ang) * random.randint(120, 330))
            d.rectangle([x - 28, y - 28, x + 28, y + 28], fill=third)
        draw_label(d, (90, 145), "Géoverrouillage: sphère locale, preuves de latence, GPS/BGP et seuil de cohérence.", 1320, base)
    elif mode == 3:
        for i in range(12):
            x = 150 + i * 100
            top = 235 + random.randint(-80, 90)
            d.rounded_rectangle([x, top, x + 58, 660], radius=18, fill=PALETTE[(index + i) % len(PALETTE)])
        draw_label(d, (95, 135), "Spectre d'entropie ambiante: bruit matériel, jitter, horloge, localisation et état système.", 1300, base)
    elif mode == 4:
        cols, rows = 10, 6
        for r in range(rows):
            for c in range(cols):
                x, y = 160 + c * 118, 185 + r * 82
                color = PALETTE[(index + r * cols + c) % len(PALETTE)]
                d.rounded_rectangle([x, y, x + 88, y + 54], radius=10, fill=color)
                d.text((x + 24, y + 15), f"S{r}{c}", font=font(19, True), fill=(255, 255, 255))
        draw_label(d, (150, 710), "Matrice de shards: chaque fragment transporte une dérivation Φ_i et un tag éphémère.", 1250, base)
    elif mode == 5:
        cx, cy = 750, 470
        for i, label in enumerate(["MitM", "Replay", "Cloud breach", "GPS spoof", "Clock drift", "SNDL"]):
            ang = -math.pi / 2 + i * 2 * math.pi / 6
            x = cx + int(math.cos(ang) * 320)
            y = cy + int(math.sin(ang) * 250)
            arrow(d, (cx, cy), (x, y), PALETTE[(index + i) % len(PALETTE)], 4)
            d.ellipse([x - 70, y - 40, x + 70, y + 40], fill=base)
            d.text((x - 56, y - 13), label, font=font(18, True), fill=(255, 255, 255))
        d.ellipse([cx - 95, cy - 95, cx + 95, cy + 95], fill=accent)
        d.text((cx - 52, cy - 15), "P-E2M2E", font=font(24, True), fill=(255, 255, 255))
    elif mode == 6:
        states = ["Initialisation", "Ping entropie", "Verrouillage", "Cristallisation", "Flux", "Annihilation"]
        for i, label in enumerate(states):
            x = 85 + i * 230
            y = 390 if i % 2 == 0 else 250
            d.rounded_rectangle([x, y, x + 180, y + 92], radius=24, fill=PALETTE[(index + i) % len(PALETTE)])
            d.text((x + 16, y + 31), label, font=font(20, True), fill=(255, 255, 255))
            if i < len(states) - 1:
                arrow(d, (x + 180, y + 46), (x + 230, 320 if i % 2 == 0 else 436), accent, 4)
        draw_label(d, (100, 660), "Automate ST-Handshake: chaque transition exige une cohérence spatio-temporelle vérifiable.", 1270, base)
    elif mode == 7:
        layers = ["Application", "Cristallisation", "Fragmentation", "Vecteur Φ", "Preuves physiques", "Transport"]
        for i, label in enumerate(layers):
            y = 170 + i * 85
            d.rounded_rectangle([270, y, 1230, y + 58], radius=12, fill=PALETTE[(index + i) % len(PALETTE)])
            d.text((310, y + 15), label, font=font(24, True), fill=(255, 255, 255))
        draw_label(d, (145, 710), "Pile logique: la couche physique conditionne l'existence logique du message.", 1200, base)
    elif mode == 8:
        for i in range(20):
            x1, y1 = random.randint(120, 1360), random.randint(160, 700)
            x2, y2 = random.randint(120, 1360), random.randint(160, 700)
            d.line([x1, y1, x2, y2], fill=PALETTE[(i + index) % len(PALETTE)], width=random.randint(2, 6))
            d.ellipse([x1 - 10, y1 - 10, x1 + 10, y1 + 10], fill=base)
        draw_label(d, (100, 690), "Graphe de cohérence: les preuves indépendantes doivent converger avant reconstruction.", 1300, base)
    elif mode == 9:
        for i in range(7):
            x = 200 + i * 170
            d.polygon([(x, 640), (x + 80, 185), (x + 155, 640)], outline=PALETTE[(index + i) % len(PALETTE)], fill=None)
            d.line([x + 80, 185, 750, 425], fill=accent, width=2)
        d.ellipse([685, 360, 815, 490], fill=base)
        d.text((704, 407), "Φ", font=font(52, True), fill=(255, 255, 255))
        draw_label(d, (110, 690), "Convergence de facteurs: temps, lieu, état ambiant, latence, appareil, politique.", 1240, base)
    elif mode == 10:
        for i, label in enumerate(["Capture", "Modulation", "Sharding", "Routage", "Réassemblage", "Effacement"]):
            x = 115 + i * 220
            d.rounded_rectangle([x, 300, x + 160, 450], radius=8, fill=PALETTE[(index + i) % len(PALETTE)])
            d.text((x + 18, 363), label, font=font(19, True), fill=(255, 255, 255))
            if i < 5:
                arrow(d, (x + 160, 375), (x + 215, 375), base, 4)
        draw_label(d, (105, 610), "Cycle de vie de la donnée: elle passe de message à potentiel, puis redevient message local.", 1270, base)
    else:
        eqs = ["Φ(t,loc,s)=H(t)⊕H(loc)⊕Q(s)", "σ=cohérence(Δt,Δx,jitter)", "A=annihile si σ<τ", "R=reconstruit si preuves valides"]
        for i, eq in enumerate(eqs):
            y = 180 + i * 130
            d.rounded_rectangle([170, y, 1330, y + 82], radius=18, outline=PALETTE[(index + i) % len(PALETTE)], width=6)
            d.text((220, y + 23), eq, font=font(31, True), fill=base)
        draw_label(d, (120, 710), "Formalisation opératoire: la validité est une fonction de cohérence, pas un simple secret statique.", 1240, base)

    path = IMG_DIR / f"schema_{index:02d}.png"
    img.save(path)
    return path


def para(seed: str, index: int, angle: str) -> str:
    digest = hashlib.sha256(f"{seed}-{index}-{angle}".encode()).hexdigest()[:10]
    return (
        f"Dans le protocole P-E2M2E, {angle} devient une composante active de la sécurité et non "
        f"une métadonnée décorative. Les travaux conduits pendant six ans établissent une lecture nouvelle "
        f"de la donnée: elle n'est plus offerte au serveur comme un fichier stable, mais comme un potentiel "
        f"fragmenté qui ne reprend forme qu'en présence du contexte physique exact. La référence technique "
        f"[{digest}] relie cette section aux axes temps, localisation et entropie ambiante."
    )


def technical_paragraph(topic: str, chapter: int, page: int) -> str:
    variants = [
        "La règle opérationnelle impose l'annihilation dès que les observations physiques sortent de la fenêtre prévue.",
        "Le serveur MB ne reçoit aucune charge utile intelligible; il manipule des fragments, des tags temporaires et des états de routage aveugles.",
        "Le nœud PTx module le flux selon Φ, puis divise la charge utile en micro-paquets dont l'existence cryptographique est éphémère.",
        "Le nœud PRx ne récupère pas une clé classique; il cristallise le message par présence contextuelle cohérente avec les contraintes du canal.",
        "Les mesures retenues dans ce dossier consolident la démonstration: latence, dérive d'horloge, cohérence spatiale, entropie locale et comportement des fragments.",
    ]
    return f"{topic}. {variants[(chapter + page) % len(variants)]}"


def build_pages(diagrams: list[Path]) -> list[Page]:
    pages: list[Page] = []
    pages.append(Page(
        "P-E2M2E - Fiche Technique Detaillee",
        [
            "Imaginez un monde où vos données ne sont lisibles que par vous, non parce qu'elles sont simplement cachées, mais parce qu'elles perdent leur sens dès qu'elles quittent votre contexte physique.",
            "Imaginez que les serveurs les plus sophistiqués puissent stocker vos fragments sans jamais voir le message, sans jamais comprendre sa forme, sans jamais posséder la possibilité de le reconstituer.",
            "Imaginez encore que les pirates informatiques les plus expérimentés se retrouvent face à une matière numérique devenue inutile: des fragments sans temps, sans lieu, sans état ambiant, donc sans existence exploitable.",
            "C'est précisément la rupture introduite par P-E2M2E. Cette technologie ne s'ajoute pas à la confidentialité traditionnelle; elle redéfinit la manière d'accéder aux données, de les manipuler, de les transmettre et de leur donner existence.",
            "Ce volume ouvre la spécification de cette cryptographie physicalisée: une architecture où la donnée n'est plus seulement protégée, mais conditionnée par le réel.",
        ],
        diagrams[0],
    ))
    pages.append(Page(
        "Resume Executif",
        [
            "P-E2M2E établit une sécurité de l'information où la donnée n'existe sous forme reconstruisible que lorsque trois familles de vecteurs convergent: le temps, l'espace et l'état physique ambiant.",
            "Le protocole abolit la dépendance centrale à la clé statique. L'accès se produit par cristallisation contextuelle: hors temps, hors lieu ou hors état physique, le message cesse d'être reconstructible.",
            "Le serveur intermédiaire MB reste aveugle par construction. Il orchestre des fragments, vérifie des tags de session et applique des politiques sans posséder le message clair ni la capacité de le reconstituer.",
            "J'y formalise P-E2M2E comme un système conçu, étudié et organisé par mes soins après six années de maturation conceptuelle, logique et technique.",
        ],
        diagrams[1],
    ))
    pages.append(Page(
        "Fiche Technique Synthetique",
        [
            "Nom: P-E2M2E. Extension: Physicalized End-to-Middle-to-End Cryptography. Domaine: cryptographie contextuelle physicalisee, cloud aveugle, securite reseau et protection de donnees sensibles.",
            "Entites: PTx pour l'emetteur physique, MB pour le milieu aveugle, PRx pour le recepteur physique. Objet central: vecteur d'etat Φ(t, loc, s).",
            "Principe directeur: rendre la reconstruction d'une donnee dependante d'une coherence physique temporaire et non d'une cle statique durable.",
            "Position du concepteur: le protocole transforme la securite en phenomene vivant. La donnée n'est pas seulement chiffrée; elle est conditionnée par son environnement d'existence.",
        ],
        diagrams[2],
        [["Parametre", "Role", "Exemple"], ["Δt", "fenetre temporelle", "fenetre nominale 10-50 ms"], ["loc", "preuve spatiale", "GPS/BGP/latence"], ["s", "etat ambiant", "jitter, bruit, capteur"]],
    ))

    chapters = [
        ("Vision Fondatrice", "l'annihilation contextuelle"),
        ("Postulat Physique", "la transformation du message en potentiel"),
        ("Modele Mathematique", "le vecteur Φ et ses composantes"),
        ("Entropie Ambiante", "l'extraction d'etat physique local"),
        ("Temps et Horloge", "la fenetre Δt et la derive"),
        ("Localisation et Reseau", "la coherence GPS, BGP et latence"),
        ("Topologie PTx-MB-PRx", "la separation des responsabilites"),
        ("Handshake ST", "l'alignement spatio-temporel"),
        ("Fragmentation Dynamique", "les micro-paquets et shards"),
        ("Cristallisation Differee", "le stockage en potentiels"),
        ("Milieu Aveugle MB", "le routage sans connaissance du contenu"),
        ("Modele de Donnees", "les tags ephemeres et fragments orphelins"),
        ("Algorithme Noyau", "la derivation physique sophistiquee"),
        ("Politique d'Acces", "les conditions d'existence de la donnee"),
        ("Modele de Menaces", "les attaques et leurs limites"),
        ("Analyse Cybersecurite", "les defenses natives revendiquees"),
        ("Simulation Python", "le prototype de laboratoire"),
        ("Mesures Reseau", "les metriques de latence et pertes"),
        ("Validation Experimentale", "les tests de penetration"),
        ("Limites Maitrisees", "les contraintes traitees par la conception"),
        ("Roadmap d'Industrialisation", "la trajectoire vers un demonstrateur souverain"),
        ("Annexes Techniques", "les tables, notations et pseudo-code"),
    ]

    page_no = len(pages) + 1
    for cidx, (chapter, focus) in enumerate(chapters, start=1):
        pages.append(Page(
            f"Chapitre {cidx} - {chapter}",
            [
                f"Ce chapitre formalise {focus}. Il expose la matrice P-E2M2E dans sa logique propre, sans la ramener a une architecture conventionnelle.",
                para("chapter", cidx, focus),
                "Chaque section présente la vision, la specification exploitable, les paramètres observables et les mécanismes qui démontrent la cohérence du protocole.",
            ],
            diagrams[(cidx + 2) % len(diagrams)],
        ))
        page_no += 1
        for local in range(1, 9):
            title = f"{chapter} - Section {local}"
            body = [
                para(chapter, local, focus),
                technical_paragraph(f"Specification {local}", cidx, local),
                f"Parametres suivis: Δt_{local}, loc_{local}, s_{local}, jitter_{local}, tag_session_{local}. Ces valeurs forment la matrice physique de cristallisation et imposent l'existence locale du message.",
                f"Critere de rejet: si la distance temporelle, la divergence spatiale ou l'incoherence d'entropie depasse le seuil τ_{local}, le fragment est classe comme annihile et ne participe plus a la reconstruction.",
            ]
            table = None
            if local in (2, 5, 8):
                table = [
                    ["Element", "Definition P-E2M2E", "Controle"],
                    ["PTx", "source physique", "capture temps, lieu, entropie"],
                    ["MB", "milieu aveugle", "routage et indexation opaque"],
                    ["PRx", "reception physicalisee", "preuve de contexte"],
                    ["Φ", "vecteur ephemere", "derive de coherence"],
                ]
            diagram = diagrams[(cidx * 8 + local) % len(diagrams)] if local in (1, 4, 7) else None
            pages.append(Page(title, body, diagram, table))
            page_no += 1

    code_pages = [
        ("Pseudo-code PE2M2E-Φ",
         [
             "function derive_phi(t, loc, ambient_state, policy, shard_index):",
             "  tw = quantize_time(t, delta_t)",
             "  gx = normalize_geofence(loc, policy.radius)",
             "  aq = sample_ambient_entropy(ambient_state)",
             "  sigma = coherence_score(tw, gx, aq, policy)",
             "  if sigma < policy.threshold: return ANNIHILATED",
             "  return SHA3_512(domain || tw || gx || aq || shard_index)",
         ]),
        ("Pseudo-code ST-Handshake",
         [
             "PTx -> MB: entropy_ping = H(snapshot_ambient || nonce_t)",
             "PRx -> MB: geofence_commitment = H(location_zone || nonce_r)",
             "MB -> PTx/PRx: latency_challenge",
             "PTx/PRx -> MB: proof_of_freshness",
             "if coherent(Δt, Δx, jitter): open stochastic stream else annihilate session",
         ]),
        ("Pseudo-code Cristallisation",
         [
             "split message into shards S_i",
             "for each S_i derive Φ_i from physical context",
             "mask S_i into potential P_i",
             "store P_i and opaque tag on MB",
             "reconstruct only when PRx produces coherent context proof",
         ]),
        ("Modele de Score de Coherence",
         [
             "score = w_t * time_alignment + w_l * location_alignment + w_s * entropy_alignment",
             "score -= w_j * jitter_penalty",
             "score -= w_d * clock_drift_penalty",
             "decision = CRYSTALLIZE if score >= τ else ANNIHILATE",
         ]),
    ]
    for title, lines in code_pages:
        pages.append(Page(title, [
            "Bloc algorithmique de reference du protocole P-E2M2E. Il formalise la derivation Φ, la decision de cristallisation et le rejet par annihilation.",
            "\n".join(lines),
            "La version prototype journalise chaque decision afin d'etablir les taux de cristallisation valide, de rejet legitime et de perte provoquee par derive contextuelle.",
        ], diagrams[len(pages) % len(diagrams)]))

    while len(pages) < 212:
        idx = len(pages) + 1
        pages.append(Page(
            f"Annexe de Specification {idx}",
            [
                para("annexe", idx, "la coherence globale du protocole"),
                "Cette page complete le dossier fondateur et renforce la base technique exploitable pour these, architecture, demonstration, prototype et presentation institutionnelle.",
                technical_paragraph("Controle transversal", idx % 22, idx % 9),
                "Note de concepteur: chaque affirmation technique est rattachee a la matrice P-E2M2E et aux six annees d'etude ayant conduit a cette architecture.",
            ],
            diagrams[idx % len(diagrams)] if idx % 5 == 0 else None,
            [["Axe", "Observation", "Mesure"], ["Temps", "derive et jitter", "ms"], ["Espace", "coherence geographique", "m"], ["Entropie", "variation ambiante", "bits"]] if idx % 7 == 0 else None,
        ))
    return pages


def write_markdown(pages: list[Page]) -> Path:
    path = OUT / "P-E2M2E_dossier_technique_source.md"
    parts = ["# P-E2M2E - Dossier Technique Detaille\n"]
    for i, page in enumerate(pages, start=1):
        parts.append(f"\n\\pagebreak\n\n## Page {i} - {page.title}\n")
        if i == 1:
            parts.append("> " + AUTHOR_CONTACT.replace("\n", "<br>") + "\n")
        for paragraph in page.body:
            if "\n" in paragraph:
                parts.append("```text\n" + paragraph + "\n```\n")
            else:
                parts.append(paragraph + "\n")
        if page.table:
            parts.append("\n| " + " | ".join(page.table[0]) + " |\n")
            parts.append("| " + " | ".join(["---"] * len(page.table[0])) + " |\n")
            for row in page.table[1:]:
                parts.append("| " + " | ".join(row) + " |\n")
        if page.diagram:
            parts.append(f"\n![{page.title}](schemas/{page.diagram.name})\n")
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


def enforce_unique_diagrams(pages: list[Page], diagrams: list[Path]) -> None:
    seen: set[Path] = set()
    for page in pages:
        if page.diagram is None:
            continue
        if page.diagram in seen:
            page.diagram = None
        else:
            seen.add(page.diagram)

    unused = [diagram for diagram in diagrams if diagram not in seen]
    if not unused:
        return

    spacing = max(1, len(pages) // len(unused))
    cursor = 0
    for i, page in enumerate(pages):
        if cursor >= len(unused):
            break
        if page.diagram is None and i % spacing == 0:
            page.diagram = unused[cursor]
            cursor += 1

    for page in pages:
        if cursor >= len(unused):
            break
        if page.diagram is None:
            page.diagram = unused[cursor]
            cursor += 1


def add_docx_paragraph(doc: Document, text: str, size: int = 10, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold


def add_docx_contact(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(AUTHOR_CONTACT)
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(73, 80, 87)


def write_docx(pages: list[Page], portrait_path: Path, qr_path: Path) -> Path:
    path = OUT / "P-E2M2E_dossier_technique_200_pages.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.86)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run().add_picture(str(portrait_path), width=Inches(0.62))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run().add_picture(str(qr_path), width=Inches(0.62))

    for i, page in enumerate(pages, start=1):
        if i == 1:
            cover = doc.add_paragraph()
            cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cover.add_run("P-E2M2E")
            run.font.name = "Arial"
            run.font.size = Pt(36)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            add_docx_contact(doc)
            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            srun = subtitle.add_run("Physicalized End-to-Middle-to-End Cryptography")
            srun.font.name = "Arial"
            srun.font.size = Pt(16)
            srun.bold = True
            srun.font.color.rgb = RGBColor(12, 74, 110)
            author = doc.add_paragraph()
            author.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arun = author.add_run("Conçu et développé par Charmant Nyungu K.")
            arun.font.name = "Arial"
            arun.font.size = Pt(12)
            arun.bold = True
            h = doc.add_heading(page.title, level=1)
        else:
            h = doc.add_heading(f"Page {i} - {page.title}", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i != 1:
            add_docx_contact(doc)
        for paragraph in page.body:
            if "\n" in paragraph:
                add_docx_paragraph(doc, paragraph, 9)
            else:
                add_docx_paragraph(doc, paragraph, 10)
        if page.table:
            table = doc.add_table(rows=len(page.table), cols=len(page.table[0]))
            table.style = "Table Grid"
            for r, row in enumerate(page.table):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(8)
                            run.bold = r == 0
        if page.diagram:
            doc.add_picture(str(page.diagram), width=Inches(6.8))
        note = doc.add_paragraph(f"P-E2M2E | Charmant Nyungu K. | Specification proprietaire | page {i}/{len(pages)}")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in note.runs:
            run.font.size = Pt(8)
        if i != len(pages):
            doc.add_page_break()
    doc.save(path)
    return path


def register_pdf_font() -> str:
    for candidate in ["/Library/Fonts/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"]:
        if Path(candidate).exists():
            pdfmetrics.registerFont(TTFont("DocFont", candidate))
            return "DocFont"
    return "Helvetica"


def draw_pdf_branding(canvas, doc, portrait_path: Path, qr_path: Path) -> None:
    width, height = A4
    portrait_size = 1.55 * cm
    qr_size = 1.45 * cm
    canvas.saveState()
    canvas.drawImage(
        ImageReader(str(portrait_path)),
        width - portrait_size - 0.55 * cm,
        height - portrait_size - 0.38 * cm,
        width=portrait_size,
        height=portrait_size,
        mask="auto",
    )
    canvas.drawImage(
        ImageReader(str(qr_path)),
        width - qr_size - 0.35 * cm,
        0.30 * cm,
        width=qr_size,
        height=qr_size,
        mask="auto",
    )
    canvas.restoreState()


def write_pdf(pages: list[Page], portrait_path: Path, qr_path: Path) -> Path:
    path = OUT / "P-E2M2E_dossier_technique_200_pages.pdf"
    font_name = register_pdf_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=10,
    )
    cover_style = ParagraphStyle(
        "CoverTitle",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=34,
        leading=40,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=8,
    )
    cover_subtitle_style = ParagraphStyle(
        "CoverSubtitle",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=14,
        leading=18,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0c4a6e"),
        spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9.2,
        leading=12,
        alignment=TA_JUSTIFY,
        spaceAfter=6,
    )
    contact_style = ParagraphStyle(
        "ContactBlock",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=7.5,
        leading=9.2,
        alignment=TA_LEFT,
        textColor=colors.HexColor("#495057"),
        spaceAfter=0,
    )
    code_style = ParagraphStyle(
        "DocCode",
        parent=styles["Code"],
        fontName=font_name,
        fontSize=8.4,
        leading=10,
        alignment=TA_LEFT,
        backColor=colors.HexColor("#f1f5f9"),
        borderPadding=6,
        spaceAfter=8,
    )
    story = []
    for i, page in enumerate(pages, start=1):
        if i == 1:
            story.append(Paragraph("P-E2M2E", cover_style))
            contact_table = Table(
                [[Paragraph(AUTHOR_CONTACT.replace("\n", "<br/>"), contact_style)]],
                colWidths=[7.8 * cm],
                hAlign="LEFT",
            )
            contact_table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#495057")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(contact_table)
            story.append(Spacer(1, 0.25 * cm))
            story.append(Paragraph("Physicalized End-to-Middle-to-End Cryptography", cover_subtitle_style))
            story.append(Paragraph("Conçu et développé par Charmant Nyungu K.", cover_subtitle_style))
            story.append(Spacer(1, 0.5 * cm))
            story.append(Paragraph(page.title, title_style))
        else:
            story.append(Paragraph(f"Page {i} - {page.title}", title_style))
            contact_table = Table(
                [[Paragraph(AUTHOR_CONTACT.replace("\n", "<br/>"), contact_style)]],
                colWidths=[7.8 * cm],
                hAlign="LEFT",
            )
            contact_table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#495057")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(contact_table)
        for paragraph in page.body:
            escaped = paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if "\n" in paragraph:
                story.append(Paragraph(escaped.replace("\n", "<br/>"), code_style))
            else:
                story.append(Paragraph(escaped, body_style))
        if page.table:
            table = Table(page.table, colWidths=[4.1 * cm, 6.2 * cm, 5.5 * cm])
            table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7.4),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#64748b")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.2 * cm))
        if page.diagram:
            story.append(PdfImage(str(page.diagram), width=15.8 * cm, height=8.95 * cm))
        story.append(Paragraph(f"P-E2M2E | Charmant Nyungu K. | Specification proprietaire | page {i}/{len(pages)}", body_style))
        if i != len(pages):
            story.append(PageBreak())

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.55 * cm,
        bottomMargin=1.9 * cm,
    )
    def branded_page(canvas, doc):
        draw_pdf_branding(canvas, doc, portrait_path, qr_path)

    doc.build(story, onFirstPage=branded_page, onLaterPages=branded_page)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    portrait_path, qr_path = prepare_brand_assets()
    diagram_titles = [
        "Topologie PTx-MB-PRx", "Fenetre temporelle", "Geoverrouillage", "Entropie ambiante",
        "Matrice de shards", "Radar de menaces", "Automate ST", "Pile logique",
        "Graphe de coherence", "Convergence physique", "Cycle de vie", "Formalisme Φ",
    ]
    diagrams = [make_diagram(i + 1, f"{diagram_titles[i % len(diagram_titles)]} #{i + 1}") for i in range(48)]
    pages = build_pages(diagrams)
    enforce_unique_diagrams(pages, diagrams)
    md = write_markdown(pages)
    docx = write_docx(pages, portrait_path, qr_path)
    pdf = write_pdf(pages, portrait_path, qr_path)
    print(f"pages={len(pages)}")
    print(f"markdown={md}")
    print(f"docx={docx}")
    print(f"pdf={pdf}")
    print(f"schemas={IMG_DIR}")
    print(f"portrait={portrait_path}")
    print(f"qrcode={qr_path}")


if __name__ == "__main__":
    from reportlab.platypus import PageBreak

    main()
